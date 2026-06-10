from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

from config import BookingStatus, room_type_names
from services.booking_service import get_booking_people_count
from services.financial_service import calculate_booking_balance, financial_totals
from services.pricing_service import calculate_revenue
from services.paths import REPORT_DIR


async def collect_financial_rows(bookings) -> list[dict]:
    rows = []
    for booking in bookings:
        people = get_booking_people_count(booking)
        calculated = await calculate_revenue(
            booking.room_type, people, booking.date_from, booking.date_to
        )
        balance = calculate_booking_balance(booking, calculated)
        rows.append(
            {
                "id": booking.id,
                "status": booking.status,
                "guest": booking.full_name or "",
                "username": booking.username or "",
                "phone": booking.phone or "",
                "room": room_type_names.get(booking.room_type, booking.room_type),
                "date_from": booking.date_from,
                "date_to": booking.date_to,
                "nights": max((booking.date_to - booking.date_from).days, 0),
                "people": people,
                **balance,
            }
        )
    return rows


async def create_financial_excel(bookings, year: int, output_dir: str | None = None) -> str:
    rows = await collect_financial_rows(bookings)
    output = Path(output_dir or REPORT_DIR)
    output.mkdir(parents=True, exist_ok=True)
    path = output / f"financial_report_{year}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заявки"
    headers = [
        "ID", "Статус", "Гость", "Username", "Телефон", "Размещение",
        "Заезд", "Выезд", "Ночей", "Гостей", "Итого", "Внесено",
        "Возвраты", "Чистая оплата", "Остаток", "Скидки", "Допуслуги",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F80ED")
    for row in rows:
        sheet.append([
            row["id"], row["status"], row["guest"], row["username"], row["phone"],
            row["room"], row["date_from"], row["date_to"], row["nights"], row["people"],
            row["total"], row["paid"], row["refunds"], row["net_paid"],
            row["remaining"], row["discount"], row["extra_services"],
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 34)

    summary = workbook.create_sheet("Итоги")
    totals = financial_totals(bookings, [row["base_total"] for row in rows])
    summary.append(["Показатель", "Сумма"])
    labels = [
        ("Начислено", "billed"), ("Внесено", "paid"), ("Возвраты", "refunds"),
        ("Чистая оплата", "net_paid"), ("Остаток", "remaining"),
        ("Скидки", "discounts"), ("Допуслуги", "extra_services"),
    ]
    for label, key in labels:
        summary.append([label, totals[key]])
    chart = BarChart()
    chart.title = f"Финансы {year}"
    chart.add_data(Reference(summary, min_col=2, min_row=1, max_row=len(labels) + 1), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=len(labels) + 1))
    summary.add_chart(chart, "D2")
    workbook.save(path)
    return str(path)


async def create_financial_pdf(bookings, year: int, output_dir: str | None = None) -> str:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    rows = await collect_financial_rows(bookings)
    output = Path(output_dir or REPORT_DIR)
    output.mkdir(parents=True, exist_ok=True)
    path = output / f"financial_report_{year}.pdf"
    totals = financial_totals(bookings, [row["base_total"] for row in rows])
    labels = ["Начислено", "Внесено", "Возвраты", "Чистая оплата", "Остаток", "Скидки", "Допуслуги"]
    values = [totals[key] for key in ("billed", "paid", "refunds", "net_paid", "remaining", "discounts", "extra_services")]
    fig, ax = plt.subplots(figsize=(11.7, 8.3))
    ax.bar(labels, values, color=["#2f80ed", "#27ae60", "#eb5757", "#219653", "#f2994a", "#9b51e0", "#56ccf2"])
    ax.set_title(f"Финансовый отчет за {year}")
    ax.set_ylabel("Рубли")
    ax.tick_params(axis="x", rotation=25)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return str(path)


def create_contacts_excel(users, output_dir: str | None = None) -> str:
    output = Path(output_dir or REPORT_DIR)
    output.mkdir(parents=True, exist_ok=True)
    path = output / f"contacts_{date.today().isoformat()}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Контакты"
    sheet.append(["Telegram ID", "Username", "Дата регистрации", "Правила приняты"])
    for user in users:
        sheet.append([user.user_id, user.username or "", user.created_at, "Да" if user.rules_accepted else "Нет"])
    sheet.freeze_panes = "A2"
    workbook.save(path)
    return str(path)
